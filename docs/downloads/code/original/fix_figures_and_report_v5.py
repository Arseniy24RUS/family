from pathlib import Path
import shutil, zipfile, math, textwrap
import numpy as np
import pandas as pd
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib import patches, patheffects as pe
from matplotlib.lines import Line2D
from matplotlib.gridspec import GridSpec
import geopandas as gpd
from shapely import wkt
from shapely.ops import unary_union
import networkx as nx
from PIL import Image, ImageStat
from docx import Document
from docx.shared import Cm
from docx.oxml.ns import qn

ROOT = Path('/mnt/data/np_family_expertise_update_v5')
if ROOT.exists():
    shutil.rmtree(ROOT)
(ROOT/'code').mkdir(parents=True)
(ROOT/'figures_fixed').mkdir(parents=True)
(ROOT/'logs').mkdir(parents=True)
(ROOT/'report').mkdir(parents=True)

DATA_DIR = Path('/mnt/data/analytics_v3/np_family_expertise_artifacts_v3/data')
ORIG_DOC = Path('/mnt/data/np_family_full_expertise_report.docx')
BASE_DOC = Path('/mnt/data/np_family_full_expertise_report_v4.docx')
OUT_DOC = ROOT/'report'/'np_family_full_expertise_report_v5.docx'

# Use a Times-like serif font available in the environment.
plt.rcParams.update({
    'font.family': 'Liberation Serif',
    'axes.unicode_minus': False,
    'font.size': 10,
    'axes.labelsize': 11,
    'xtick.labelsize': 9,
    'ytick.labelsize': 9,
    'legend.fontsize': 8,
    'figure.facecolor': 'white',
    'axes.facecolor': 'white',
    'savefig.facecolor': 'white',
    'savefig.bbox': 'tight',
})

PALETTE = {
    'blue': '#5B84B1',
    'teal': '#6FA4A0',
    'green': '#7AA874',
    'red': '#C86464',
    'orange': '#D79B49',
    'purple': '#8E78B3',
    'gray': '#B9B9B9',
    'dark': '#404040',
    'light': '#F5F6F7',
}
CLUSTER_COLORS = ['#5B84B1','#7AA874','#C86464','#D79B49','#8E78B3','#6FA4A0','#B98E6B','#A7A7A7']
COMMUNITY_COLORS = ['#5B84B1','#7AA874','#C86464','#D79B49','#8E78B3','#6FA4A0','#B98E6B','#A7A7A7','#E07B91','#7F8C8D']

SHORT = {
    '2.14.Я.1':'Бедность многодетных семей',
    '2.14.Я.2':'Суммарный коэффициент рождаемости',
    '2.14.Я.3':'СКР третьих и последующих детей',
    '2.14.Я.4':'Долговременный уход',
    '2.14.Я.5':'Центры аудиовизуального контента',
    '2.14.Я.6':'Удовлетворённость культурой',
    '2.14.Я1.1':'Ежегодная семейная выплата',
    '2.14.Я1.2':'Проактивные меры поддержки',
    '2.14.Я1.3':'Капремонт детсадов',
    '2.14.Я1.4':'Комнаты матери и ребёнка в вузах',
    '2.14.Я1.5':'Поддержка студенческих семей',
    '2.14.Я2.1':'Социальный контракт',
    '2.14.Я2.2':'Группы продлённого дня',
    '2.14.Я2.3':'Прирост многодетных семей',
    '2.14.Я2.4':'Регионы с темпами выше среднероссийских',
    '2.14.Я3.1':'Охват диспансеризацией',
    '2.14.Я3.2':'Помощь женщинам в сельской местности',
    '2.14.Я3.3':'Младенческая смертность',
    '2.14.Я3.4':'Диспансерное наблюдение детей',
    '2.14.Я3.5':'Помощь беременным',
    '2.14.Я4.1':'Активное долголетие',
    '2.14.Я5.1':'Семья как главная ценность',
    '2.14.Я5.2':'Посещения культуры',
}

# ---------------- helpers ----------------
def wrap_text(s, width):
    return '\n'.join(textwrap.wrap(str(s), width=width, break_long_words=False))


def draw_number(ax, x, y, s, size=7.2, weight='bold'):
    ax.text(x, y, str(s), fontsize=size, fontweight=weight, ha='center', va='center', color='black',
            path_effects=[pe.withStroke(linewidth=2.2, foreground='white')], zorder=6)


def draw_legend_grid(ax, entries, ncols=3, fontsize=7.0, wrap_width=34, bullet_colors=None):
    """Entries are strings. Optional bullet_colors aligned with entries."""
    ax.axis('off')
    entries = [wrap_text(e, wrap_width) for e in entries]
    n = len(entries)
    nrows = math.ceil(n / ncols)
    col_width = 1 / ncols
    # line height in axes units
    line_h = 0.028 if fontsize >= 7 else 0.025
    row_gap = 0.010
    y = 0.98
    # column-major fill
    cols = []
    colors_cols = []
    for c in range(ncols):
        start = c * nrows
        end = min((c + 1) * nrows, n)
        cols.append(entries[start:end])
        if bullet_colors is None:
            colors_cols.append([None] * len(cols[-1]))
        else:
            colors_cols.append(bullet_colors[start:end])
    max_rows = max(len(c) for c in cols)
    for r in range(max_rows):
        row_items = []
        row_colors = []
        max_lines = 1
        for c in range(ncols):
            if r < len(cols[c]):
                item = cols[c][r]
                row_items.append(item)
                row_colors.append(colors_cols[c][r])
                max_lines = max(max_lines, item.count('\n') + 1)
            else:
                row_items.append(None)
                row_colors.append(None)
        for c, item in enumerate(row_items):
            if item is None:
                continue
            x = c * col_width + 0.012
            if row_colors[c] is not None:
                ax.add_patch(patches.Rectangle((x, y - line_h*0.8), 0.012, 0.012, transform=ax.transAxes,
                                               facecolor=row_colors[c], edgecolor='black', linewidth=0.3))
                text_x = x + 0.018
            else:
                text_x = x
            ax.text(text_x, y, item, transform=ax.transAxes, ha='left', va='top', fontsize=fontsize, color='black')
        y -= max_lines * line_h + row_gap
    
# ---------------- geography ----------------
territory_ref = pd.read_csv(DATA_DIR/'territory_reference_89.csv')
poly_raw = pd.read_csv('/mnt/data/Карта 89 субъектов РФ с населением.csv')
poly_raw['geometry'] = poly_raw['WKT'].apply(wkt.loads).apply(unary_union)
poly = gpd.GeoDataFrame(poly_raw[['Name_full','geometry']], geometry='geometry', crs='EPSG:4326')
name_map = {
    'Республика Адыгея': 'Республика Адыгея (Адыгея)',
    'Республика Татарстан': 'Республика Татарстан (Татарстан)',
    'Республика Северная Осетия - Алания': 'Республика Северная Осетия – Алания',
    'Кемеровская область - Кузбасс': 'Кемеровская область – Кузбасс',
    'Ханты-Мансийский автономный округ - Югра': 'Ханты-Мансийский автономный округ – Югра',
    'Москва': 'Город Москва',
    'Санкт-Петербург': 'Город Санкт-Петербург',
    'Севастополь': 'Город Севастополь',
}
territory_ref['polygon_name'] = territory_ref['region_name'].map(name_map).fillna(territory_ref['region_name'])
geo = territory_ref.merge(poly, left_on='polygon_name', right_on='Name_full', how='left')
geo = gpd.GeoDataFrame(geo, geometry='geometry', crs='EPSG:4326')
DISPLAY_CRS = '+proj=laea +lat_0=60 +lon_0=100 +datum=WGS84 +units=m +no_defs'
display_geo = geo.to_crs(DISPLAY_CRS)
display_geo['geometry'] = display_geo.geometry.buffer(0).simplify(4000, preserve_topology=True)
cent = display_geo.geometry.centroid
display_geo['cx'] = cent.x
display_geo['cy'] = cent.y
BOUNDS = display_geo.total_bounds


def setup_map_ax(ax):
    xmin, ymin, xmax, ymax = BOUNDS
    xpad = (xmax - xmin) * 0.02
    ypad = (ymax - ymin) * 0.03
    ax.set_xlim(xmin - xpad, xmax + xpad)
    ax.set_ylim(ymin - ypad, ymax + ypad)
    ax.set_aspect('equal')
    ax.set_axis_off()

# ---------------- figures ----------------
def fig1_logic_model(path):
    fig = plt.figure(figsize=(13.2, 4.9), dpi=220)
    ax = fig.add_axes([0.01, 0.01, 0.98, 0.98])
    ax.set_xlim(0, 1)
    ax.set_ylim(0, 1)
    ax.axis('off')

    xs = [0.01, 0.205, 0.40, 0.595, 0.79]
    w = 0.18
    y = 0.18
    h = 0.72
    colors = ['#E9F1F8','#EAF5EC','#FBF2E2','#F7E9E6','#EEE8F7']
    titles = [
        'Ресурсы',
        'Мероприятия',
        'Непосредственные\nрезультаты',
        'Промежуточные\nрезультаты',
        'Итоговые\nэффекты'
    ]
    texts = [
        'Финансирование\nНормативная база\nФедеральные и региональные\nисполнители\nИнформационные системы\nМежведомственная координация',
        'Выплаты и льготы\nРазвитие инфраструктуры\nПоддержка репродуктивного\nздоровья\nДолговременный уход\nКультурные инициативы',
        'Получатели мер поддержки\nОтремонтированные объекты\nОхват диспансеризацией\nОснащённые консультации\nСозданные сервисы\nРост доступности услуг',
        'Снижение бедности\nСнижение барьеров рождения\nУлучшение здоровья\nУкрепление устойчивости семей\nРост доступности ухода',
        'Рост реализованной\nрождаемости\nСнижение младенческой\nсмертности\nПовышение качества жизни\nсемей\nУстойчивое демографическое\nразвитие'
    ]
    for x, c, title, txt in zip(xs, colors, titles, texts):
        box = patches.FancyBboxPatch((x, y), w, h, boxstyle='round,pad=0.009,rounding_size=0.016',
                                     facecolor=c, edgecolor='#666666', linewidth=1.1)
        ax.add_patch(box)
        ax.text(x + w/2, y + h - 0.035, title, ha='center', va='top', fontsize=11.5, fontweight='bold')
        ax.text(x + w/2, y + h - 0.125, txt, ha='center', va='top', fontsize=9.1, linespacing=1.26)
    for i in range(len(xs) - 1):
        ax.annotate('', xy=(xs[i+1] - 0.007, y + h/2), xytext=(xs[i] + w + 0.007, y + h/2),
                    arrowprops=dict(arrowstyle='-|>', lw=1.9, color='#666666'))
    band = patches.FancyBboxPatch((0.06, 0.045), 0.88, 0.065, boxstyle='round,pad=0.007,rounding_size=0.01',
                                  facecolor='#F7F7F7', edgecolor='#888888', linewidth=0.8)
    ax.add_patch(band)
    ax.text(0.50, 0.078,
            'Экспертная логика: проверка согласованности ресурсов, мероприятий, системы показателей, территориальных различий и конечного демографического эффекта',
            ha='center', va='center', fontsize=9.1)
    fig.savefig(path)
    plt.close(fig)


def fig14_pca(path):
    df = pd.read_csv(DATA_DIR/'region_clusters.csv').copy().sort_values('region_name').reset_index(drop=True)
    df['num'] = np.arange(1, len(df) + 1)
    cluster_ids = sorted(df['cluster_id'].dropna().astype(int).unique())
    color_map = {cid: CLUSTER_COLORS[i % len(CLUSTER_COLORS)] for i, cid in enumerate(cluster_ids)}

    fig = plt.figure(figsize=(13.8, 14.6), dpi=220)
    gs = GridSpec(2, 1, height_ratios=[3.0, 2.3], hspace=0.08)
    ax = fig.add_subplot(gs[0])
    for cid, part in df.groupby('cluster_id'):
        cid = int(cid)
        ax.scatter(part['pca1'], part['pca2'], s=55, color=color_map[cid], edgecolors='black', linewidths=0.35,
                   alpha=0.92, label=f'Кластер {cid}', zorder=3)
    for _, r in df.iterrows():
        draw_number(ax, r['pca1'], r['pca2'], int(r['num']), size=6.7)
    ax.axhline(0, color='#888888', lw=0.8)
    ax.axvline(0, color='#888888', lw=0.8)
    ax.grid(alpha=0.22)
    ax.set_xlabel('Первая главная компонента')
    ax.set_ylabel('Вторая главная компонента')
    ax.legend(loc='upper center', bbox_to_anchor=(0.5, 1.02), ncol=min(4, len(cluster_ids)), frameon=False)
    ax.text(0.01, 0.01, 'Номера точек соответствуют субъектам РФ в легенде ниже', transform=ax.transAxes, fontsize=8)

    ax2 = fig.add_subplot(gs[1])
    entries = [f"{int(n)} — {name}" for n, name in zip(df['num'], df['region_name'])]
    draw_legend_grid(ax2, entries, ncols=4, fontsize=7.1, wrap_width=28)

    fig.savefig(path)
    plt.close(fig)


def fig19_network(path):
    edges = pd.read_csv(DATA_DIR/'activities_indicators_matrix_long.csv').copy()
    acts = edges[['activity_id', 'activity']].drop_duplicates().sort_values('activity_id').reset_index(drop=True)
    inds = edges[['indicator_code']].drop_duplicates().sort_values('indicator_code').reset_index(drop=True)
    inds['indicator_name'] = inds['indicator_code'].map(lambda x: SHORT.get(x, x))
    acts['num'] = np.arange(1, len(acts) + 1)
    inds['num'] = np.arange(1, len(inds) + 1)

    G = nx.Graph()
    for _, r in acts.iterrows():
        G.add_node('A_' + r['activity_id'], kind='activity', label=r['activity'], num=int(r['num']))
    for _, r in inds.iterrows():
        G.add_node('I_' + r['indicator_code'], kind='indicator', label=r['indicator_name'], num=int(r['num']))
    for _, r in edges.iterrows():
        G.add_edge('A_' + r['activity_id'], 'I_' + r['indicator_code'], relationship=r['relationship_type'])

    act_nodes = ['A_' + x for x in acts['activity_id']]
    ind_nodes = ['I_' + x for x in inds['indicator_code']]
    pos = {}
    for i, n in enumerate(act_nodes, start=1):
        pos[n] = (0.18, 1 - i / (len(act_nodes) + 1))
    for i, n in enumerate(ind_nodes, start=1):
        pos[n] = (0.82, 1 - i / (len(ind_nodes) + 1))

    fig = plt.figure(figsize=(14.4, 20.0), dpi=220)
    gs = GridSpec(3, 1, height_ratios=[3.0, 1.55, 1.3], hspace=0.08)
    ax = fig.add_subplot(gs[0])
    ax.axis('off')
    rel_color = {'direct': PALETTE['blue'], 'indirect': '#C8C8C8'}
    rel_width = {'direct': 1.1, 'indirect': 0.8}
    for u, v, d in G.edges(data=True):
        ax.plot([pos[u][0], pos[v][0]], [pos[u][1], pos[v][1]], color=rel_color.get(d['relationship'], '#CCCCCC'),
                lw=rel_width.get(d['relationship'], 0.8), alpha=0.7, zorder=1)
    ax.scatter([pos[n][0] for n in act_nodes], [pos[n][1] for n in act_nodes], s=125,
               c='#B9D3EA', edgecolors='black', linewidths=0.35, zorder=3)
    ax.scatter([pos[n][0] for n in ind_nodes], [pos[n][1] for n in ind_nodes], s=140, marker='s',
               c='#E9C2C2', edgecolors='black', linewidths=0.35, zorder=3)
    for n in act_nodes + ind_nodes:
        x, y = pos[n]
        draw_number(ax, x, y, G.nodes[n]['num'], size=6.4)
    ax.text(0.12, 1.01, 'Мероприятия', transform=ax.transAxes, fontsize=11, fontweight='bold')
    ax.text(0.76, 1.01, 'Показатели', transform=ax.transAxes, fontsize=11, fontweight='bold')
    handles = [
        Line2D([0], [0], marker='o', markersize=8, markerfacecolor='#B9D3EA', markeredgecolor='black', color='none', label='Мероприятия'),
        Line2D([0], [0], marker='s', markersize=8, markerfacecolor='#E9C2C2', markeredgecolor='black', color='none', label='Показатели'),
        Line2D([0, 1], [0, 0], color=PALETTE['blue'], lw=1.1, label='Прямая связь'),
        Line2D([0, 1], [0, 0], color='#C8C8C8', lw=0.8, label='Косвенная связь'),
    ]
    ax.legend(handles=handles, ncol=4, loc='lower center', bbox_to_anchor=(0.5, -0.03), frameon=False)

    ax1 = fig.add_subplot(gs[1])
    ax1.text(0.0, 1.02, 'Нумерация мероприятий', transform=ax1.transAxes, fontsize=10, fontweight='bold')
    act_entries = [f"{int(n)} — {a}" for n, a in zip(acts['num'], acts['activity'])]
    draw_legend_grid(ax1, act_entries, ncols=2, fontsize=6.7, wrap_width=42)

    ax2 = fig.add_subplot(gs[2])
    ax2.text(0.0, 1.02, 'Нумерация показателей', transform=ax2.transAxes, fontsize=10, fontweight='bold')
    ind_entries = [f"{int(n)} — {code}: {name}" for n, code, name in zip(inds['num'], inds['indicator_code'], inds['indicator_name'])]
    draw_legend_grid(ax2, ind_entries, ncols=2, fontsize=6.8, wrap_width=44)

    fig.savefig(path)
    plt.close(fig)


def fig24_region_network(path):
    edges = pd.read_csv(DATA_DIR/'region_similarity_network_edges.csv').copy()
    comm = pd.read_csv(DATA_DIR/'region_similarity_network_communities.csv').copy().sort_values('region_name').reset_index(drop=True)
    comm['num'] = np.arange(1, len(comm) + 1)
    num_map = dict(zip(comm['region_code'], comm['num']))
    name_map = dict(zip(comm['region_code'], comm['region_name']))
    community_map = dict(zip(comm['region_code'], comm['community_id']))

    G = nx.Graph()
    for _, r in comm.iterrows():
        G.add_node(int(r['region_code']), label=r['region_name'], num=int(r['num']), community=int(r['community_id']))

    # Keep enough edges for visible network while avoiding clutter.
    threshold = float(edges['similarity_weight'].quantile(0.72))
    edges_use = edges[edges['similarity_weight'] >= threshold].copy()
    if len(edges_use) < 50:
        edges_use = edges.nlargest(70, 'similarity_weight').copy()
    for _, r in edges_use.iterrows():
        G.add_edge(int(r['region_code_1']), int(r['region_code_2']), weight=float(r['similarity_weight']))

    pos = nx.spring_layout(G, seed=42, weight='weight', k=0.42, iterations=350)
    communities = sorted(comm['community_id'].unique())
    comm_color = {cid: COMMUNITY_COLORS[i % len(COMMUNITY_COLORS)] for i, cid in enumerate(communities)}

    fig = plt.figure(figsize=(14.5, 18.5), dpi=220)
    gs = GridSpec(2, 1, height_ratios=[3.0, 2.4], hspace=0.06)
    ax = fig.add_subplot(gs[0])
    ax.axis('off')
    # edges
    max_w = edges_use['similarity_weight'].max()
    min_w = edges_use['similarity_weight'].min()
    for u, v, d in G.edges(data=True):
        lw = 0.4 + 1.2 * ((d['weight'] - min_w) / (max_w - min_w + 1e-9))
        ax.plot([pos[u][0], pos[v][0]], [pos[u][1], pos[v][1]], color='#CFCFCF', lw=lw, alpha=0.75, zorder=1)
    for cid in communities:
        nodes = [n for n in G.nodes if G.nodes[n]['community'] == cid]
        ax.scatter([pos[n][0] for n in nodes], [pos[n][1] for n in nodes], s=45,
                   c=comm_color[cid], edgecolors='black', linewidths=0.25, zorder=3, label=f'Сообщество {cid}')
    for n in G.nodes:
        draw_number(ax, pos[n][0], pos[n][1], G.nodes[n]['num'], size=6.2)
    ax.legend(loc='upper center', bbox_to_anchor=(0.5, 1.02), ncol=min(4, len(communities)), frameon=False)
    ax.text(0.01, 0.01, 'Номера точек соответствуют субъектам РФ в легенде ниже', transform=ax.transAxes, fontsize=8)

    # Full legend with community membership
    ax2 = fig.add_subplot(gs[1])
    entries = [f"{int(n)} — {name} (сообщество {int(cid)})" for n, name, cid in zip(comm['num'], comm['region_name'], comm['community_id'])]
    colors = [comm_color[int(cid)] for cid in comm['community_id']]
    draw_legend_grid(ax2, entries, ncols=4, fontsize=6.9, wrap_width=30, bullet_colors=colors)

    fig.savefig(path)
    plt.close(fig)


def choropleth(indicator_code, path, legend_label):
    regional_latest = pd.read_csv(DATA_DIR/'regional_latest_values_89_long.csv')
    sub = regional_latest[regional_latest['indicator_code'] == indicator_code].copy()
    sub = sub.sort_values(['region_code', 'year']).drop_duplicates('region_code', keep='last')
    g = display_geo.merge(sub[['region_code', 'value']], on='region_code', how='left')
    fig, ax = plt.subplots(figsize=(12.0, 8.0), dpi=220)
    g.plot(column='value', ax=ax, cmap='YlGnBu', linewidth=0.3, edgecolor='white', missing_kwds={'color': '#E3E3E3', 'edgecolor': 'white'})
    g.boundary.plot(ax=ax, color='#8A8A8A', linewidth=0.2)
    setup_map_ax(ax)
    vals = g['value'].dropna()
    sm = plt.cm.ScalarMappable(cmap='YlGnBu', norm=plt.Normalize(vmin=float(vals.min()), vmax=float(vals.max())))
    sm._A = []
    cbar = fig.colorbar(sm, ax=ax, fraction=0.026, pad=0.02)
    cbar.set_label(legend_label)
    fig.savefig(path)
    plt.close(fig)


def composite_map(path):
    df = pd.read_csv(DATA_DIR/'regional_composite_indices.csv')[['region_code', 'family_project_performance_index']]
    g = display_geo.merge(df, on='region_code', how='left')
    fig, ax = plt.subplots(figsize=(12.0, 8.0), dpi=220)
    g.plot(column='family_project_performance_index', ax=ax, cmap='YlGnBu', linewidth=0.3, edgecolor='white', missing_kwds={'color': '#E3E3E3', 'edgecolor': 'white'})
    g.boundary.plot(ax=ax, color='#8A8A8A', linewidth=0.2)
    setup_map_ax(ax)
    vals = g['family_project_performance_index'].dropna()
    sm = plt.cm.ScalarMappable(cmap='YlGnBu', norm=plt.Normalize(vmin=float(vals.min()), vmax=float(vals.max())))
    sm._A = []
    cbar = fig.colorbar(sm, ax=ax, fraction=0.026, pad=0.02)
    cbar.set_label('Интегральный индекс (0–100)')
    fig.savefig(path)
    plt.close(fig)


def lisa_map(indicator_code, path):
    lisa = pd.read_csv(DATA_DIR/'spatial_local_moran_long.csv')
    sub = lisa[lisa['indicator_code'] == indicator_code].copy().sort_values('region_code').drop_duplicates('region_code', keep='last')
    g = display_geo.merge(sub[['region_code', 'cluster_type']], on='region_code', how='left')
    colors = {'HH':'#D97B7B','LL':'#77A6D1','HL':'#E5C07B','LH':'#8FC2B0','NS':'#D9D9D9'}
    labels = {'HH':'Высокое–высокое','LL':'Низкое–низкое','HL':'Высокое–низкое','LH':'Низкое–высокое','NS':'Незначимо / нет данных'}
    fig, ax = plt.subplots(figsize=(12.0, 8.0), dpi=220)
    for ct in ['HH','LL','HL','LH','NS']:
        part = g[g['cluster_type'].fillna('NS') == ct]
        part.plot(ax=ax, color=colors[ct], edgecolor='white', linewidth=0.3)
    g.boundary.plot(ax=ax, color='#8A8A8A', linewidth=0.2)
    setup_map_ax(ax)
    handles = [patches.Patch(facecolor=colors[k], edgecolor='white', label=labels[k]) for k in ['HH','LL','HL','LH','NS']]
    ax.legend(handles=handles, loc='lower left', frameon=True, facecolor='white')
    fig.savefig(path)
    plt.close(fig)

# ---------------- render figures ----------------
FIXED = {
    1: ROOT/'figures_fixed'/'figure_01_logic_model_fixed_v5.png',
    14: ROOT/'figures_fixed'/'figure_14_pca_clusters_fixed_v5.png',
    16: ROOT/'figures_fixed'/'figure_16_composite_map_fixed_v5.png',
    19: ROOT/'figures_fixed'/'figure_19_activity_indicator_network_fixed_v5.png',
    24: ROOT/'figures_fixed'/'figure_24_region_similarity_network_fixed_v5.png',
    28: ROOT/'figures_fixed'/'figure_28_rural_women_map_fixed_v5.png',
    30: ROOT/'figures_fixed'/'figure_30_children_followup_map_fixed_v5.png',
    31: ROOT/'figures_fixed'/'figure_31_pregnant_help_map_fixed_v5.png',
    35: ROOT/'figures_fixed'/'figure_35_lisa_tfr_fixed_v5.png',
    36: ROOT/'figures_fixed'/'figure_36_lisa_tfr3_fixed_v5.png',
}
fig1_logic_model(FIXED[1])
fig14_pca(FIXED[14])
composite_map(FIXED[16])
fig19_network(FIXED[19])
fig24_region_network(FIXED[24])
choropleth('2.14.Я3.2', FIXED[28], 'Процент')
choropleth('2.14.Я3.4', FIXED[30], 'Процент')
choropleth('2.14.Я3.5', FIXED[31], 'Процент')
lisa_map('2.14.Я.2', FIXED[35])
lisa_map('2.14.Я.3', FIXED[36])

# QC
qc_rows = []
for num, p in FIXED.items():
    img = Image.open(p).convert('L')
    stat = ImageStat.Stat(img)
    # estimate non-white bbox
    arr = np.array(img)
    mask = arr < 250
    if mask.any():
        ys, xs = np.where(mask)
        bbox = [int(xs.min()), int(ys.min()), int(xs.max()), int(ys.max())]
        fill_ratio = float(mask.mean())
    else:
        bbox = None
        fill_ratio = 0.0
    qc_rows.append({'figure_number': num, 'filename': p.name, 'size': f'{img.size[0]}x{img.size[1]}', 'mean_gray': stat.mean[0], 'std_gray': stat.stddev[0], 'fill_ratio': fill_ratio, 'bbox': str(bbox)})
pd.DataFrame(qc_rows).sort_values('figure_number').to_csv(ROOT/'logs'/'fixed_figures_qc_v5.csv', index=False)

# ---------------- patch docx ----------------
base_doc = Document(str(BASE_DOC))
orig_doc = Document(str(ORIG_DOC))

def get_image_rids(doc):
    ids = []
    for p in doc.paragraphs:
        if p._element.xpath('.//w:drawing'):
            for b in p._element.xpath('.//a:blip'):
                ids.append(b.get(qn('r:embed')))
    return ids

base_ids = get_image_rids(base_doc)
orig_ids = get_image_rids(orig_doc)
assert len(base_ids) == 36 and len(orig_ids) == 36

# replace selected images only
for fig_no, img_path in FIXED.items():
    rid = base_ids[fig_no - 1]
    base_doc.part.related_parts[rid]._blob = img_path.read_bytes()

# sizing: keep original sizes for figures 28, 30, 31 to avoid horizontal squashing.
orig_shapes = list(orig_doc.inline_shapes)
base_shapes = list(base_doc.inline_shapes)
for idx, shape in enumerate(base_shapes, start=1):
    if idx in [28, 30, 31]:
        shape.width = orig_shapes[idx - 1].width
        shape.height = orig_shapes[idx - 1].height
    elif idx in FIXED:
        img = Image.open(FIXED[idx])
        if idx in [1, 16, 28, 30, 31, 35, 36]:
            target_width_cm = 15.5
        else:
            target_width_cm = 15.0
        target_height_cm = target_width_cm * img.size[1] / img.size[0]
        shape.width = Cm(target_width_cm)
        shape.height = Cm(target_height_cm)

# save doc
base_doc.save(str(OUT_DOC))

# README and manifest
readme = ROOT/'README.md'
readme.write_text(
    'Пакет содержит исправления версии 5.\n\n'
    'Исправлены рисунки 1, 14, 16, 19, 24, 28, 30, 31, 35 и 36.\n'
    'Для рисунков 28, 30 и 31 восстановлены корректные пропорции карт в документе и заново построены изображения в полигональной проекции.\n'
    'Для рисунков 14 и 24 сохранены цветовые группы, добавлены номера поверх цветных точек и полные легенды по всем субъектам РФ.\n'
    'Для рисунка 19 сохранены цветовые различия узлов, исправлена легенда и устранены наложения текста.\n'
    'Для рисунка 1 полностью переработана компоновка с уменьшением пустых полей и устранением выхода текста за границы.\n',
    encoding='utf-8'
)
files = []
for p in ROOT.rglob('*'):
    if p.is_file():
        files.append({'path': str(p.relative_to(ROOT)), 'size_bytes': p.stat().st_size})
pd.DataFrame(files).sort_values('path').to_csv(ROOT/'logs'/'package_manifest_v5.csv', index=False)

shutil.copy2('/mnt/data/fix_v5_script.py', ROOT/'code'/'fix_figures_and_report_v5.py')
zip_path = Path('/mnt/data/np_family_expertise_update_v5_package.zip')
if zip_path.exists():
    zip_path.unlink()
shutil.make_archive('/mnt/data/np_family_expertise_update_v5_package', 'zip', ROOT)
print('saved', OUT_DOC)
print('saved', zip_path)
